import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import pandas as pd

doc = docx.Document()

# Set page margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

# Header Title
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('PUNE INSTITUTE OF COMPUTER TECHNOLOGY, PUNE - 411043\nDEPARTMENT OF COMPUTER ENGINEERING\nSUBJECT & LOAD ALLOCATION TO FACULTY (S.Y. 2026-27 SEM-I)')
run.bold = True
run.font.size = Pt(12)
run.font.name = 'Calibri'

raw_records = [
    # (Teacher Name, Theory Subject, Theory Class, Theory Hrs, Practical Subj, Practical Class/Batch, Practical Hrs)
    ('Dr. A. G. Phakatkar', 'DS', 'SE-2', 3, 'DSL', 'SE-2 (H2)', 4),
    ('Dr. A. S. Ghotkar', 'DM', 'SE-1', 3, '', '', 0),
    ('Dr. G. V. Kale', 'DM', 'SE-3', 3, '', '', 0),
    ('Dr. K. C. Waghmare', 'DS', 'SE-4', 3, 'DSL', 'SE-1 (F1)', 4),
    ('Dr. K. C. Waghmare', '', '', 0, 'DSL', 'SE-3 (F3)', 4),
    ('Dr. P. R. Patil', 'DS', 'SE-3', 3, '', '', 0),
    ('Dr. S. A. Joshi', 'DS', 'SE-1', 3, 'DSL', 'SE-4 (G4)', 4),
    ('Dr. S. N. Girme', 'COA', 'SE-1', 3, 'COAL', 'SE-1 (E1)', 2),
    ('Dr. S. N. Girme', '', '', 0, 'COAL', 'SE-2 (E2)', 2),
    ('Dr. S. N. Girme', '', '', 0, 'COAL', 'SE-3 (F3)', 2),
    ('Dr. S. S. Sonawane', 'DM', 'SE-2', 3, '', '', 0),
    ('Dr. Sheetal Sonawane', '', '', 0, 'DSL', 'SE-4 (F4)', 4),
    ('Prof. A. A. Chandorkar', 'MDM', 'SE-3', 2, 'MDM TUT', 'SE-1 (E1)', 1),
    ('Prof. A. A. Chandorkar', '', '', 0, 'MDM TUT', 'SE-2 (E2)', 1),
    ('Prof. A. A. Chandorkar', '', '', 0, 'MDM TUT', 'SE-2 (F2)', 1),
    ('Prof. A. A. Chandorkar', '', '', 0, 'MDM TUT', 'SE-3 (G3)', 1),
    ('Prof. A. A. Chandorkar', '', '', 0, 'MDM TUT', 'SE-3 (H3)', 1),
    ('Prof. A. A. Chandorkar', '', '', 0, 'MDM TUT', 'SE-4 (F4)', 1),
    ('Prof. A. D. Bundele', '', '', 0, 'COAL', 'SE-3 (G3)', 2),
    ('Prof. A. D. Bundele', '', '', 0, 'COAL', 'SE-4 (E4)', 2),
    ('Prof. A. D. Bundele', '', '', 0, 'COAL', 'SE-4 (H4)', 2),
    ('Prof. B. P. Masram', '', '', 0, 'CEP', 'SE-1 (F1)', 2),
    ('Prof. D. D. Raigar', 'COA', 'SE-2', 3, 'COAL', 'SE-1 (G1)', 2),
    ('Prof. D. D. Raigar', '', '', 0, 'CEP', 'SE-1 (H1)', 2),
    ('Prof. D. D. Raigar', '', '', 0, 'COAL', 'SE-2 (F2)', 2),
    ('Prof. D. D. Raigar', '', '', 0, 'COAL', 'SE-2 (G2)', 2),
    ('Prof. D. D. Raigar', '', '', 0, 'COAL', 'SE-4 (G4)', 2),
    ('Prof. D. D. Raigar', '', '', 0, 'CEP', 'SE-4 (F4)', 2),
    ('Prof. H. Khan', '', '', 0, 'PDCR', 'SE-1 (F1)', 2),
    ('Prof. H. Khan', '', '', 0, 'PDCR', 'SE-4 (F4)', 2),
    ('Prof. Jayshree Mahajan', 'UHV', 'SE-2', 2, 'FLS', 'SE-2 (F2)', 2),
    ('Prof. Jayshree Mahajan', 'UHV', 'SE-4', 2, 'PDCR', 'SE-3 (F3)', 2),
    ('Prof. Jayshree Mahajan', '', '', 0, 'PDCR', 'SE-3 (G3)', 2),
    ('Prof. Jayshree Mahajan', '', '', 0, 'FLS', 'SE-3 (G3)', 2),
    ('Prof. Kopal Gangrade', '', '', 0, 'CEP', 'SE-1 (G1)', 2),
    ('Prof. Kopal Gangrade', '', '', 0, 'PDCR', 'SE-2 (F2)', 2),
    ('Prof. Kopal Gangrade', '', '', 0, 'PDCR', 'SE-2 (H2)', 2),
    ('Prof. Kopal Gangrade', '', '', 0, 'CEP', 'SE-2 (G2)', 2),
    ('Prof. Kopal Gangrade', '', '', 0, 'PDCR', 'SE-4 (H4)', 2),
    ('Prof. M. Patil', 'COA', 'SE-4', 3, 'COAL', 'SE-2 (H2)', 2),
    ('Prof. M. Patil', '', '', 0, 'COAL', 'SE-3 (E3)', 2),
    ('Prof. M. Patil', '', '', 0, 'CEP', 'SE-3 (H3)', 2),
    ('Prof. M. Patil', '', '', 0, 'CEP', 'SE-4 (E4)', 2),
    ('Prof. M. Patil', '', '', 0, 'CEP', 'SE-4 (G4)', 2),
    ('Prof. M. R. Jansari', '', '', 0, 'DSL', 'SE-1 (H1)', 4),
    ('Prof. M. R. Jansari', '', '', 0, 'PDCR', 'SE-1 (H1)', 2),
    ('Prof. M. S. Chavan', '', '', 0, 'DSL', 'SE-4 (E4)', 4),
    ('Prof. M. S. Wakode', '', '', 0, 'DSL', 'SE-3 (E3)', 4),
    ('Prof. M. S. Wakode', '', '', 0, 'DSL', 'SE-4 (H4)', 4),
    ('Prof. M. V. Mane', 'MDM', 'SE-4', 2, 'MDM TUT', 'SE-1 (H1)', 1),
    ('Prof. M. V. Mane', '', '', 0, 'CEP', 'SE-2 (F2)', 2),
    ('Prof. M. V. Mane', '', '', 0, 'CEP', 'SE-2 (H2)', 2),
    ('Prof. M. V. Mane', '', '', 0, 'MDM TUT', 'SE-2 (G2)', 1),
    ('Prof. M. V. Mane', '', '', 0, 'MDM TUT', 'SE-2 (H2)', 1),
    ('Prof. M. V. Mane', '', '', 0, 'CEP', 'SE-3 (F3)', 2),
    ('Prof. M. V. Mane', '', '', 0, 'MDM TUT', 'SE-3 (E3)', 1),
    ('Prof. M. V. Mane', '', '', 0, 'MDM TUT', 'SE-4 (G4)', 1),
    ('Prof. M. V. Mane', '', '', 0, 'MDM TUT', 'SE-4 (H4)', 1),
    ('Prof. M. V. Mane', '', '', 0, 'CEP', 'SE-4 (H4)', 2),
    ('Prof. P. A. Chavan', 'EEFM', 'SE-4', 1, 'FLS', 'SE-1 (G1)', 2),
    ('Prof. P. A. Chavan', '', '', 0, 'FLS', 'SE-4 (E4)', 2),
    ('Prof. P. A. Chavan', '', '', 0, 'FLS', 'SE-4 (H4)', 2),
    ('Prof. P. A. Jain', 'DM', 'SE-4', 3, 'DSL', 'SE-1 (G1)', 4),
    ('Prof. P. A. Jain', '', '', 0, 'DSL', 'SE-2 (F2)', 4),
    ('Prof. P. J. Jambhulkar', '', '', 0, 'DSL', 'SE-2 (E2)', 4),
    ('Prof. P. J. Jambhulkar', '', '', 0, 'DSL', 'SE-3 (G3)', 4),
    ('Prof. P. J. Jambhulkar', '', '', 0, 'PDCR', 'SE-3 (H3)', 2),
    ('Prof. P. J. Jambhulkar', '', '', 0, 'PDCR', 'SE-4 (G4)', 2),
    ('Prof. P. P. Joshi', 'MDM', 'SE-2', 2, 'PDCR', 'SE-2 (E2)', 2),
    ('Prof. P. P. Joshi', '', '', 0, 'CEP', 'SE-3 (G3)', 2),
    ('Prof. P. R. Navghare', 'MDM', 'SE-1', 2, 'PDCR', 'SE-1 (G1)', 2),
    ('Prof. P. R. Navghare', '', '', 0, 'PDCR', 'SE-3 (E3)', 2),
    ('Prof. P. R. Navghare', '', '', 0, 'FLS', 'SE-3 (F3)', 2),
    ('Prof. P. R. Navghare', '', '', 0, 'FLS', 'SE-3 (H3)', 2),
    ('Prof. P. R. Navghare', '', '', 0, 'FLS', 'SE-4 (G4)', 2),
    ('Prof. P. S. Shanane', 'EEFM', 'SE-1', 1, '', '', 0),
    ('Prof. P. S. Shanane', 'EEFM', 'SE-2', 1, '', '', 0),
    ('Prof. P. S. Shanane', 'EEFM', 'SE-3', 1, '', '', 0),
    ('Prof. R. R. Jadhav', '', '', 0, 'MDM TUT', 'SE-1 (F1)', 1),
    ('Prof. R. R. Jadhav', '', '', 0, 'MDM TUT', 'SE-1 (G1)', 1),
    ('Prof. R. R. Jadhav', '', '', 0, 'FLS', 'SE-1 (F1)', 2),
    ('Prof. R. R. Jadhav', '', '', 0, 'FLS', 'SE-2 (E2)', 2),
    ('Prof. R. R. Jadhav', '', '', 0, 'FLS', 'SE-2 (G2)', 2),
    ('Prof. R. R. Jadhav', '', '', 0, 'MDM TUT', 'SE-3 (F3)', 1),
    ('Prof. R. R. Jadhav', '', '', 0, 'MDM TUT', 'SE-4 (E4)', 1),
    ('Prof. R. R. Jadhav', '', '', 0, 'FLS', 'SE-4 (F4)', 2),
    ('Prof. R. S. Paswan', 'COA', 'SE-3', 3, 'COAL', 'SE-1 (F1)', 2),
    ('Prof. R. S. Paswan', '', '', 0, 'COAL', 'SE-1 (H1)', 2),
    ('Prof. R. S. Paswan', '', '', 0, 'COAL', 'SE-3 (H3)', 2),
    ('Prof. R. S. Paswan', '', '', 0, 'COAL', 'SE-4 (F4)', 2),
    ('Prof. Rutuja A. Kulkarni', '', '', 0, 'DSL', 'SE-1 (E1)', 4),
    ('Prof. Rutuja A. Kulkarni', '', '', 0, 'CEP', 'SE-1 (E1)', 2),
    ('Prof. Rutuja A. Kulkarni', '', '', 0, 'CEP', 'SE-2 (E2)', 2),
    ('Prof. Rutuja A. Kulkarni', '', '', 0, 'CEP', 'SE-3 (E3)', 2),
    ('Prof. S. P. Shintre', '', '', 0, 'DSL', 'SE-3 (H3)', 4),
    ('Prof. S. Shah', '', '', 0, 'PDCR', 'SE-2 (G2)', 2),
    ('Prof. S. Shah', '', '', 0, 'PDCR', 'SE-4 (E4)', 2),
    ('Prof. S. W. Jadhav', 'UHV', 'SE-1', 2, 'PDCR', 'SE-1 (E1)', 2),
    ('Prof. S. W. Jadhav', 'UHV', 'SE-3', 2, 'FLS', 'SE-1 (E1)', 2),
    ('Prof. S. W. Jadhav', '', '', 0, 'FLS', 'SE-1 (H1)', 2),
    ('Prof. S. W. Jadhav', '', '', 0, 'FLS', 'SE-2 (H2)', 2),
    ('Prof. S. W. Jadhav', '', '', 0, 'FLS', 'SE-3 (E3)', 2),
    ('Prof. S. W. Jadhav', '', '', 0, 'FLS', 'SE-4 (H4)', 2),
    ('Prof. Y. A. Handge', '', '', 0, 'DSL', 'SE-2 (G2)', 4)
]

# Calculate teacher total load
teacher_totals = {}
for name, th_s, th_c, th_h, pr_s, pr_c, pr_h in raw_records:
    if name not in teacher_totals:
        teacher_totals[name] = 0
    teacher_totals[name] += th_h + pr_h

# Create Table in DOCX
table = doc.add_table(rows=1, cols=8)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

hdr_cells = table.rows[0].cells
headers = ['Name of the Teacher', 'Theory Subject', 'Class', 'Theory Load', 'Practical / Lab Subject', 'Batch / Class', 'Lab Load', 'Total Load']
for i, header_text in enumerate(headers):
    cell = hdr_cells[i]
    cell.text = header_text
    cell.paragraphs[0].runs[0].bold = True
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Background shading
    shd_xml = parse_xml(r'<w:shd {} w:fill="1F4E78"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shd_xml)
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    cell.paragraphs[0].runs[0].font.size = Pt(9.5)

# Populate rows
last_teacher = None
for row_data in raw_records:
    t_name, th_s, th_c, th_h, pr_s, pr_c, pr_h = row_data
    row_cells = table.add_row().cells
    
    # Teacher Name
    row_cells[0].text = t_name if t_name != last_teacher else ''
    row_cells[1].text = th_s
    row_cells[2].text = th_c
    row_cells[3].text = str(th_h) if th_h > 0 else ''
    row_cells[4].text = pr_s
    row_cells[5].text = pr_c
    row_cells[6].text = str(pr_h) if pr_h > 0 else ''
    row_cells[7].text = str(teacher_totals[t_name]) if t_name != last_teacher else ''
    
    last_teacher = t_name

# Save DOCX
doc_path = 'Load allocation-26-27 SEM-1 - SE Computer Engg.docx'
doc.save(doc_path)
print(f'Successfully created {doc_path}')

# Create CSV and Excel
df_data = []
for row_data in raw_records:
    t_name, th_s, th_c, th_h, pr_s, pr_c, pr_h = row_data
    df_data.append({
        'teacher_name': t_name,
        'theory_subject': th_s,
        'theory_class': th_c,
        'theory_hours': th_h,
        'practical_subject': pr_s,
        'practical_batch': pr_c,
        'practical_hours': pr_h,
        'total_load': teacher_totals[t_name]
    })

df = pd.DataFrame(df_data)
csv_path = 'load_allocation_se_2026_27.csv'
excel_path = 'Load allocation-26-27 SEM-1 - SE Computer Engg.xlsx'
df.to_csv(csv_path, index=False)
df.to_excel(excel_path, index=False)

print(f'Successfully created {csv_path}')
print(f'Successfully created {excel_path}')
